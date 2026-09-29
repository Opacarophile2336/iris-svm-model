"""Export service for Iris prediction history — CSV and Excel (XLSX) generation.

Reads raw records from SQL Server Predictions and formats them into publication-ready files.
"""
from __future__ import annotations

import csv
import io
from typing import Any, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


EXPORT_HEADERS = [
    ("Prediction ID", "prediction_id"),
    ("User ID", "user_id"),
    ("Username", "username"),
    ("Created At", "created_at"),
    ("Sepal Length (cm)", "sepal_length"),
    ("Sepal Width (cm)", "sepal_width"),
    ("Petal Length (cm)", "petal_length"),
    ("Petal Width (cm)", "petal_width"),
    ("Predicted Species", "predicted_species"),
    ("Probability Setosa (%)", "probability_setosa"),
    ("Probability Versicolor (%)", "probability_versicolor"),
    ("Probability Virginica (%)", "probability_virginica"),
    ("Model Name", "model_name"),
    ("Kernel / Type", "kernel_display"),
]


def _format_cell_value(key: str, val: Any) -> Any:
    if val is None:
        return ""
    if key in ("prediction_id", "user_id"):
        try:
            return int(val)
        except (ValueError, TypeError):
            return str(val)
    if key in ("probability_setosa", "probability_versicolor", "probability_virginica"):
        # Format as percentage rounded to 2 decimals
        num = float(val)
        if num <= 1.0:
            num *= 100.0
        return round(num, 2)
    if key in ("sepal_length", "sepal_width", "petal_length", "petal_width"):
        return round(float(val), 2)
    return str(val)


def generate_history_csv(records: List[dict[str, Any]], is_admin: bool = False) -> str:
    """Generate RFC 4180 compliant CSV string from prediction history records."""
    headers_to_use = EXPORT_HEADERS if is_admin else [h for h in EXPORT_HEADERS if h[1] != "username"]

    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\r\n")

    # Header row
    writer.writerow([h[0] for h in headers_to_use])

    # Data rows
    for r in records:
        row = [_format_cell_value(h[1], r.get(h[1])) for h in headers_to_use]
        writer.writerow(row)

    return output.getvalue()


def generate_history_excel(records: List[dict[str, Any]], is_admin: bool = False) -> bytes:
    """Generate a professionally styled XLSX workbook in memory using openpyxl."""
    headers_to_use = EXPORT_HEADERS if is_admin else [h for h in EXPORT_HEADERS if h[1] != "username"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prediction History"

    # Styling definitions
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Segoe UI", size=10)
    data_align_center = Alignment(horizontal="center", vertical="center")
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    # Write Header
    ws.row_dimensions[1].height = 26
    for col_idx, (header_label, _) in enumerate(headers_to_use, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Write Data Rows
    for row_idx, r in enumerate(records, 2):
        ws.row_dimensions[row_idx].height = 20
        use_alt = (row_idx % 2 == 1)

        for col_idx, (_, field_key) in enumerate(headers_to_use, 1):
            val = _format_cell_value(field_key, r.get(field_key))
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if use_alt:
                cell.fill = alt_fill

            # Formatting & alignment based on type
            if field_key in ("prediction_id", "user_id"):
                cell.alignment = data_align_center
            elif field_key in ("sepal_length", "sepal_width", "petal_length", "petal_width",
                              "probability_setosa", "probability_versicolor", "probability_virginica"):
                cell.alignment = data_align_right
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            elif field_key == "created_at":
                cell.alignment = data_align_center
            else:
                cell.alignment = data_align_left

    # Auto-adjust column widths with safety margin
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
